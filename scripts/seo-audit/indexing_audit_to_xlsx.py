import json
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill
from openpyxl.utils import get_column_letter

DIR = r"D:\personal\Propertify\scripts\seo-audit"
rows = json.load(open(DIR + r"\indexing-audit.json", encoding="utf-8"))
before = json.load(open(DIR + r"\googlebot-crawl-before.json", encoding="utf-8"))
after = json.load(open(DIR + r"\googlebot-crawl.json", encoding="utf-8"))

FONT = "Arial"
HDR = Font(name=FONT, bold=True, color="FFFFFF")
HDR_FILL = PatternFill("solid", fgColor="2F5597")
BODY = Font(name=FONT, size=10)
BOLD = Font(name=FONT, size=10, bold=True)
TITLE = Font(name=FONT, size=14, bold=True, color="2F5597")
RED = PatternFill("solid", fgColor="F8CBAD")
YEL = PatternFill("solid", fgColor="FFF2CC")
GRN = PatternFill("solid", fgColor="E2EFDA")

ACTIONS = [
    ("WAS ORPHAN", "Fixed (Oct 7): now linked from the listing pages"),
    ("WAS 1 LINK", "Fixed (Oct 7): now linked from listings + site navigation"),
    ("no English content", "Write the English version of this page's article (currently only Arabic text)"),
    ("no Arabic content", "Write the Arabic version of this page's article"),
    ("thin content", "Expand the page text - aim for 300+ words of real, specific content"),
    ("duplicate title", "Give this page its own SEO title - another page uses the exact same one"),
    ("title in wrong language", "Fix the name/title field: it's in the other language"),
    ("no price", "Fill 'Price value (number)' in the admin panel"),
    ("no meta description", "Low priority: add a custom SEO description (a fallback is used now)"),
]


def actions_for(issues):
    return "\n".join(f"• {a}" for k, a in ACTIONS if k in issues) or "Nothing - this page is in good shape"


def header(ws, cols, widths):
    for i, c in enumerate(cols, 1):
        cell = ws.cell(row=1, column=i, value=c)
        cell.font, cell.fill = HDR, HDR_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(i)].width = widths.get(c, 14)
    ws.freeze_panes = "A2"


def stats(rs):
    sm = [r for r in rs if r["in_sitemap"]]
    ms = sorted(r["ms"] for r in sm if isinstance(r["ms"], (int, float)))
    return {
        "orphans": sum(1 for r in sm if not r["reached_by_links"]),
        "one_link": sum(1 for r in sm if r["inbound_links"] == 1),
        "within2": sum(1 for r in sm if r["reached_by_links"] and r["depth"] <= 2),
        "maxdepth": max(r["depth"] for r in sm if r["reached_by_links"]),
        "median_ms": ms[len(ms) // 2],
    }


wb = Workbook()

# ---------------- Summary ----------------
ws = wb.active
ws.title = "Summary"
ws.column_dimensions["A"].width = 46
ws.column_dimensions["B"].width = 18
ws.column_dimensions["C"].width = 18
ws.column_dimensions["D"].width = 60
r = 1
ws.cell(row=r, column=1, value="Aqar Factory — Indexing Audit (crawled as Googlebot, Oct 7 2026)").font = TITLE
r += 2
ws.cell(row=r, column=1, value="WHAT GOOGLEBOT SEES — before vs after today's fix").font = BOLD
r += 1
for i, h in enumerate(["Metric", "Before", "After", "Why it matters"], 1):
    c = ws.cell(row=r, column=i, value=h)
    c.font, c.fill = HDR, HDR_FILL
b, a = stats(before), stats(after)
metrics = [
    ("Sitemap URLs with NO internal link (orphans)", b["orphans"], a["orphans"], "Google only knew them from the sitemap -> lowest crawl priority = 'Discovered - currently not indexed'"),
    ("Sitemap URLs with only 1 internal link", b["one_link"], a["one_link"], "Weak importance signal"),
    ("Pages within 2 clicks of the homepage", b["within2"], a["within2"], "Shallow pages get crawled first"),
    ("Deepest page (clicks from homepage)", b["maxdepth"], a["maxdepth"], ""),
    ("Median crawler response time (ms, uncached)", b["median_ms"], a["median_ms"], "Repeat crawls are now cached: ~100 ms (was never cached before)"),
    ("Pages returning an error / wrong canonical / noindex", 0, 0, "Nothing was ever technically BLOCKING indexing"),
]
for m in metrics:
    r += 1
    for i, v in enumerate(m, 1):
        c = ws.cell(row=r, column=i, value=v)
        c.font = BODY
        c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.cell(row=r, column=3).fill = GRN if m[2] != m[1] else PatternFill()

r += 2
ws.cell(row=r, column=1, value="REMAINING CONTENT WORK (needs a writer / the client — see 'Content To-Do' tab)").font = BOLD
counts = {k: sum(1 for x in rows if k in x["issues"]) for k, _ in ACTIONS[2:]}
for k, act in ACTIONS[2:]:
    r += 1
    ws.cell(row=r, column=1, value=act).font = BODY
    ws.cell(row=r, column=1).alignment = Alignment(wrap_text=True)
    ws.cell(row=r, column=2, value=counts[k]).font = BOLD
r += 1
ws.cell(row=r, column=1, value="Pages with no remaining issue").font = BODY
ws.cell(row=r, column=2, value=sum(1 for x in rows if x["remaining_content_issues"] == 0)).font = BOLD
ws.cell(row=r, column=2).fill = GRN

r += 2
ws.cell(row=r, column=1, value="WHAT TO DO IN GOOGLE SEARCH CONSOLE NOW").font = BOLD
steps = [
    "1. Sitemaps -> resubmit https://www.aqar-factory.com/sitemap.xml so Google re-reads it with the new link structure.",
    "2. Leave the running 'Discovered - currently not indexed' validation alone - do not restart it.",
    "3. URL Inspection -> 'Request indexing' for your 10 most important projects (quota is ~10/day).",
    "4. Expect movement over 2-6 weeks, not days: Google has to re-crawl the listing pages, find the new links, then crawl each page.",
    "5. Off-site: claim/verify a Google Business Profile with the office address, and get links from developer partners and social profiles - a 2-month-old domain needs outside links for Google to raise its crawl budget.",
]
for s in steps:
    r += 1
    c = ws.cell(row=r, column=1, value=s)
    c.font = BODY
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=4)
    c.alignment = Alignment(wrap_text=True)
    ws.row_dimensions[r].height = 30

# ---------------- All URLs ----------------
ws2 = wb.create_sheet("All URLs (522)")
cols = ["url", "type", "lang", "name", "linked_before", "inbound_before", "inbound_after", "depth_before", "depth_after",
        "status", "canonical_ok", "body_words", "faq_pairs", "issues", "what_to_do"]
widths = {"url": 55, "name": 34, "issues": 44, "what_to_do": 60, "title": 40}
header(ws2, cols, widths)
for i, x in enumerate(sorted(rows, key=lambda x: (-x["remaining_content_issues"], x["type"], x["url"])), 2):
    x = dict(x, what_to_do=actions_for(x["issues"]))
    for j, c in enumerate(cols, 1):
        cell = ws2.cell(row=i, column=j, value=x.get(c, ""))
        cell.font = BODY
        cell.alignment = Alignment(vertical="top", wrap_text=c in ("issues", "what_to_do", "name"))
    ws2.cell(row=i, column=5).fill = RED if x["linked_before"] == "NO" else PatternFill()
    ws2.cell(row=i, column=14).fill = GRN if x["remaining_content_issues"] == 0 else YEL
ws2.auto_filter.ref = ws2.dimensions

# ---------------- Content To-Do ----------------
ws3 = wb.create_sheet("Content To-Do")
cols3 = ["url", "type", "lang", "name", "body_words", "issues", "what_to_do", "done"]
header(ws3, cols3, {"url": 55, "name": 34, "issues": 44, "what_to_do": 60, "done": 10})
todo = sorted([x for x in rows if x["remaining_content_issues"] > 0], key=lambda x: (-x["remaining_content_issues"], x["type"], x["url"]))
for i, x in enumerate(todo, 2):
    x = dict(x, what_to_do=actions_for(x["issues"]), done="")
    for j, c in enumerate(cols3, 1):
        cell = ws3.cell(row=i, column=j, value=x.get(c, ""))
        cell.font = BODY
        cell.alignment = Alignment(vertical="top", wrap_text=c in ("issues", "what_to_do", "name"))
    ws3.cell(row=i, column=8).fill = YEL
ws3.auto_filter.ref = ws3.dimensions

out = DIR + r"\indexing-audit.xlsx"
wb.save(out)
print(f"wrote {out}: {len(rows)} URLs, {len(todo)} with content work")
